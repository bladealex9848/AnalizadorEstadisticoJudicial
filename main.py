import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from pathlib import Path
import tempfile
import os
import openpyxl
from openpyxl.styles import Alignment, Border, Side
import base64
from io import BytesIO
import re
import glob
from datetime import datetime
from openpyxl.utils.dataframe import dataframe_to_rows
import warnings
import zipfile

# Configuración de la página
st.set_page_config(page_title="AnalizadorEstadisticoJudicial", page_icon="📊", layout="wide")

# Inicializar archivo de log
def initialize_log():
    log_path = Path("log.txt")
    if not log_path.exists():
        with open(log_path, "w") as log_file:
            log_file.write("Registro para AnalizadorEstadisticoJudicial Web\n\n")
    return log_path

def log_message(message, add_space=False, log_file="log.txt"):
    """
    Función que registra los mensajes en un archivo log.
    """
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    log_entry = f"[{timestamp}] {message}"
    with open(log_file, "a") as f:
        if add_space:
            f.write("\n")
        f.write(log_entry + "\n")

def create_download_link(url, label):
    return f'<a href="{url}" target="_blank" class="btn-download">{label}</a>'

def show_sidebar_resources():
    st.sidebar.title("Recursos Adicionales")
    with st.sidebar.expander("Ver Recursos Adicionales", expanded=False):        
        st.markdown(create_download_link("https://enki.care/AnalizadorEstadisticoJudicialGuiaUsuario", 'Guía de Usuario'), unsafe_allow_html=True)

    st.sidebar.title("Marco Normativo")
    with st.sidebar.expander("Ver Marco Normativo", expanded=False):
        st.markdown(create_download_link("https://enki.care/PSAA16-10618", 'ACUERDO PSAA16-10618'), unsafe_allow_html=True)
        st.markdown(create_download_link("https://enki.care/SIERJU-Formularios-e-Instructivos", 'SIERJU Formularios e Instructivos'), unsafe_allow_html=True)
        
    st.sidebar.title("Descargar Versión Portable")
    with st.sidebar.expander("Ver Versión Portable", expanded=False):
        st.markdown(create_download_link("https://enki.care/AnalizadorEstadisticoJudicial_v1.1.zip", 'Versión 1.1'), unsafe_allow_html=True)

    st.sidebar.markdown("""
    <style>
    .btn-download {
        display: inline-block;
        padding: 0.5em 1em;
        color: #ffffff !important;
        background-color: #0066cc;
        border-radius: 5px;
        text-decoration: none;
        font-weight: bold;
        text-align: center;
        margin-bottom: 5px;
    }
    .btn-download:hover {
        background-color: #0056b3;
        color: #ffffff !important;
        text-decoration: none;
    }
    </style>
    """, unsafe_allow_html=True)

def get_binary_file_downloader_html(bin_file, file_label='File'):
    with open(bin_file, 'rb') as f:
        data = f.read()
    bin_str = base64.b64encode(data).decode()
    href = f'<a href="data:application/octet-stream;base64,{bin_str}" download="{os.path.basename(bin_file)}" class="btn-download">Descargar {file_label}</a>'
    return href

def sort_key_func(file_name):
    """
    Función que devuelve una clave de orden para los archivos.
    """
    order = ["Primer", "Segundo", "Tercer", "Cuarto"]
    match = re.match(r"(\w+ Trimestre)(_?(\d)?)", Path(file_name).name)
    if match:
        name, _, number = match.groups()
        return (order.index(name.split(' ')[0]), int(number) if number else 0)
    return (0, 0)

def sorted_files(files):
    """
    Función que ordena los archivos según su orden de trimestre y número.
    """
    return sorted(files, key=sort_key_func)

def create_folder_structure(temp_dir):
    """
    Función que crea estructura de carpetas para guardar resultados.
    """
    log_message("Creando estructura de carpetas para guardar resultados.")
    subfolder = os.path.join(temp_dir, 'Consolidado', datetime.now().strftime('%Y-%m-%d_%H-%M-%S'))
    os.makedirs(subfolder, exist_ok=True)
    return subfolder

def process_excel_files(excel_files, subfolder, log_file):
    """
    Procesa los archivos Excel y devuelve los datos consolidados.
    """
    log_message("Procesando archivos Excel.", log_file=log_file)
    all_sheets_data = {}

    # Ordenar los archivos por nombre para asegurar que se procesan en el orden correcto
    excel_files.sort(key=sort_key_func)

    for file in excel_files:
        file_path = Path(file)
        if "Trimestre" not in file_path.name:
            continue
            
        log_message(f"Procesando archivo: {file_path.name}", log_file=log_file)
        st.info(f"Procesando archivo: {file_path.name}")
        
        try:
            if file_path.suffix.lower() == '.xls':
                xls = pd.ExcelFile(file, engine='xlrd')
            else:
                xls = pd.ExcelFile(file, engine='openpyxl')
                
            result_file = os.path.join(subfolder, file_path.stem + '_results.xlsx')
            writer = openpyxl.Workbook()
            writer.remove(writer.active)

            process_sheets(xls, file_path.name, all_sheets_data, writer, subfolder, log_file)

            try:
                writer.save(result_file)
                log_message(f"Archivo de resultados guardado: {result_file}", log_file=log_file)
                st.success(f"Resultados guardados para {file_path.name}")
            except Exception as e:
                log_message(f"Error al guardar el archivo de resultados: {str(e)}", log_file=log_file)
                st.error(f"Error al guardar resultados para {file_path.name}: {str(e)}")
                
        except Exception as e:
            log_message(f"Error al procesar el archivo {file_path.name}: {str(e)}", log_file=log_file)
            st.error(f"Error al procesar {file_path.name}: {str(e)}")

    return all_sheets_data

def process_sheets(xls, file_name, all_sheets_data, writer, subfolder, log_file):
    """
    Procesa cada hoja del archivo Excel.
    """
    log_message(f"Procesando hojas del archivo: {file_name}", log_file=log_file)
    
    for sheet in xls.sheet_names:
        log_message(f"Procesando hoja: {sheet}", log_file=log_file)
        
        try:
            data = pd.read_excel(xls, sheet_name=sheet, header=None)
            
            # Reemplazar NaN por None para evitar problemas
            data = data.replace({np.nan: None})
            
            if len(data) >= 20 and 'Total' in data[0].values:
                process_rows(data, file_name, all_sheets_data, writer, sheet, log_file)
                st.write(f"Hoja '{sheet}' procesada correctamente.")
            else:
                log_message(f"La hoja {sheet} del archivo {file_name} no cumple con las condiciones necesarias", log_file=log_file)
                st.warning(f"Hoja '{sheet}' en {file_name} no cumple con el formato esperado.")
                
        except Exception as e:
            log_message(f"Error al leer la hoja {sheet} del archivo {file_name}: {str(e)}", log_file=log_file)
            st.error(f"Error al procesar hoja '{sheet}': {str(e)}")

def process_rows(data, file_name, all_sheets_data, writer, sheet, log_file):
    """
    Procesa las filas de datos de una hoja.
    """
    log_message(f"Procesando datos de la hoja {sheet} del archivo {file_name}", log_file=log_file)
    
    header_rows = data.iloc[:19].values.tolist()
    row_20_titles = [''] + data.iloc[19, 1:].tolist()
    
    # Buscar la fila 'Total'
    total_indices = data.index[data[0] == 'Total'].tolist()
    if not total_indices:
        log_message(f"No se encontró fila 'Total' en la hoja {sheet} del archivo {file_name}", log_file=log_file)
        st.warning(f"No se encontró fila 'Total' en la hoja {sheet}")
        return
        
    total_row_index = total_indices[0]
    total_row_values = ['Total'] + data.iloc[total_row_index, 1:].tolist()

    try:
        last_column_index = len(row_20_titles)
    except:
        last_column_index = len(row_20_titles)

    header_rows = [row[:last_column_index] for row in header_rows]
    row_20_titles = row_20_titles[:last_column_index]
    total_row_values = total_row_values[:last_column_index]

    results_df = pd.DataFrame(header_rows + [row_20_titles, total_row_values])

    ws = writer.create_sheet(title=sheet)
    for r_idx, row in enumerate(dataframe_to_rows(results_df, index=False, header=False)):
        ws.append(row)

    # Ajustar el texto para los títulos de la fila 19
    for cell in ws[19]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=Side(style='thin'),
                             right=Side(style='thin'),
                             top=Side(style='thin'),
                             bottom=Side(style='thin'))

    # Ajustar el texto para las celdas desde la fila 20 hasta el final
    for row in ws.iter_rows(min_row=20, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=Side(style='thin'),
                                 right=Side(style='thin'),
                                 top=Side(style='thin'),
                                 bottom=Side(style='thin'))

    # Combinar celdas contiguas con el mismo valor en la fila 19
    prev_val = ws.cell(row=19, column=1).value
    merge_start = None
    for col in range(2, ws.max_column + 1):
        current_val = ws.cell(row=19, column=col).value
        if current_val == prev_val:
            if not merge_start:
                merge_start = col - 1
            if col == ws.max_column:
                ws.merge_cells(start_row=19, start_column=merge_start, end_row=19, end_column=col)
        else:
            if merge_start:
                ws.merge_cells(start_row=19, start_column=merge_start, end_row=19, end_column=col - 1)
                merge_start = None
        prev_val = current_val

    if sheet not in all_sheets_data:
        all_sheets_data[sheet] = []
    
    all_sheets_data[sheet].append(total_row_values + [file_name])

def consolidate_data(data):
    """
    Función que consolida los datos por trimestre.
    """
    consolidated_data = []
    trimesters = ["Primer Trimestre", "Segundo Trimestre", "Tercer Trimestre", "Cuarto Trimestre"]

    for trimester in trimesters:
        rows = [row for row in data if trimester in row[-1]]
        if rows:
            # Asumiendo que todas las filas tienen la misma estructura
            # Sumar los valores numéricos para cada columna
            consolidated_row = ['Total']
            for i in range(1, len(rows[0]) - 1):
                try:
                    # Intentar sumar los valores numéricos
                    consolidated_row.append(sum(float(row[i]) if row[i] is not None else 0 for row in rows))
                except (ValueError, TypeError):
                    # Si no se puede convertir a número, usar el último valor no nulo
                    values = [row[i] for row in rows if row[i] is not None]
                    consolidated_row.append(values[-1] if values else None)
                    
            consolidated_row.append(trimester)
            consolidated_data.append(consolidated_row)

    return consolidated_data

def create_consolidated_file(all_sheets_data, subfolder, log_file):
    """
    Función que crea el archivo consolidado con todos los resultados.
    """
    log_message("Iniciando la creación del archivo consolidado.", log_file=log_file)
    consolidated_writer = openpyxl.Workbook()
    consolidated_writer.remove(consolidated_writer.active)

    border = Border(left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin'))

    for sheet, data in all_sheets_data.items():
        ws = consolidated_writer.create_sheet(title=sheet)

        # Verificar si necesitamos una segunda tabla
        has_multiple_parts = any("_" in str(item[-1]) for item in data)
        
        # Primera tabla con datos individuales
        for row in sorted(data, key=lambda x: sort_key_func(x[-1])):
            ws.append(row)

        # Si hay múltiples partes, agregamos la segunda tabla
        if has_multiple_parts:
            # Agregar un espacio entre las dos tablas
            ws.append([])

            # Segunda tabla
            consolidated_rows = consolidate_data(data)
            for row in consolidated_rows:
                ws.append(row)

        for row in ws.rows:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True)
                cell.border = border

    consolidated_file = os.path.join(subfolder, 'Consolidado.xlsx')
    try:
        consolidated_writer.save(consolidated_file)
        log_message(f"Archivo consolidado creado exitosamente en {consolidated_file}.", log_file=log_file)
        st.success("Archivo consolidado creado exitosamente.")
        return consolidated_file
    except Exception as e:
        log_message(f"Error al guardar el archivo consolidado: {str(e)}", log_file=log_file)
        st.error(f"Error al crear archivo consolidado: {str(e)}")
        return None

def create_zip_file(subfolder, log_file_path):
    """
    Crea un archivo ZIP con todos los resultados y el log.
    """
    zip_path = os.path.join(subfolder, 'Resultados_AnalizadorEstadisticoJudicial.zip')
    
    with zipfile.ZipFile(zip_path, 'w') as zipf:
        # Agregar todos los archivos del subfolder
        for root, _, files in os.walk(subfolder):
            for file in files:
                if file.endswith('.zip'):  # Evitar incluir el propio ZIP
                    continue
                file_path = os.path.join(root, file)
                arcname = os.path.relpath(file_path, subfolder)
                zipf.write(file_path, arcname)
        
        # Agregar el archivo de log
        if os.path.exists(log_file_path):
            zipf.write(log_file_path, "log.txt")
    
    return zip_path

def load_sample_dataset():
    """
    Carga un dataset de muestra para demostración.
    """
    # Dataset de muestra para casos de derecho de familia en Colombia
    sample_data = {
        "Casos de Familia": [
            ["Total", 560, 500, 60, 69, "Primer Trimestre"],
            ["Total", 560, 545, 75, 70, "Segundo Trimestre"],
            ["Total", 620, 585, 110, 69.6, "Tercer Trimestre"],
            ["Total", 615, 625, 100, 68.8, "Cuarto Trimestre"]
        ],
        "Eficiencia Judicial": [
            ["Total", 0.89, 1.12, 0.12, "Primer Trimestre"],
            ["Total", 0.97, 1.14, 0.14, "Segundo Trimestre"],
            ["Total", 0.94, 1.19, 0.19, "Tercer Trimestre"],
            ["Total", 1.02, 1.16, 0.16, "Cuarto Trimestre"]
        ]
    }
    return sample_data

def show_summary(all_sheets_data):
    """
    Muestra un resumen de los datos procesados.
    """
    st.header("Resumen de Datos")
    
    if not all_sheets_data:
        st.warning("No hay datos para mostrar. Por favor, carga y procesa los archivos Excel.")
        return
        
    for sheet, data in all_sheets_data.items():
        st.subheader(f"Hoja: {sheet}")
        
        # Convertir datos a DataFrame para mejor visualización
        try:
            df = pd.DataFrame(data, columns=["Concepto"] + [f"Valor {i+1}" for i in range(len(data[0])-2)] + ["Trimestre"])
            st.dataframe(df)
        except Exception as e:
            st.error(f"Error al mostrar datos de la hoja {sheet}: {str(e)}")
            st.write("Datos en bruto:", data)

def show_details(all_sheets_data):
    """
    Muestra detalles por trimestre.
    """
    st.header("Detalles por Trimestre")
    
    if not all_sheets_data:
        st.warning("No hay datos para mostrar. Por favor, carga y procesa los archivos Excel.")
        return
    
    trimester = st.selectbox("Selecciona un trimestre", 
                             ["Primer Trimestre", "Segundo Trimestre", "Tercer Trimestre", "Cuarto Trimestre"])
    
    for sheet, data in all_sheets_data.items():
        # Filtrar datos por trimestre seleccionado
        trimester_data = [row for row in data if trimester in row[-1]]
        
        if trimester_data:
            st.subheader(f"{sheet} - {trimester}")
            try:
                df = pd.DataFrame(trimester_data, columns=["Concepto"] + [f"Valor {i+1}" for i in range(len(trimester_data[0])-2)] + ["Trimestre"])
                st.dataframe(df)
            except Exception as e:
                st.error(f"Error al mostrar detalles de la hoja {sheet} para {trimester}: {str(e)}")
                st.write("Datos en bruto:", trimester_data)

def show_charts(all_sheets_data):
    """
    Muestra gráficos de los datos procesados.
    """
    st.header("Visualización de Datos")
    
    if not all_sheets_data:
        st.warning("No hay datos para visualizar. Por favor, carga y procesa los archivos Excel.")
        return
    
    sheet = st.selectbox("Selecciona una hoja", list(all_sheets_data.keys()))
    
    if all_sheets_data[sheet]:
        try:
            # Transformar datos para gráficos
            data = all_sheets_data[sheet]
            df = pd.DataFrame(data)
            
            # La última columna es el trimestre
            df.columns = ["Concepto"] + [f"Valor_{i}" for i in range(len(df.columns)-2)] + ["Trimestre"]
            
            # Convertir columnas numéricas
            for col in df.columns:
                if col not in ["Concepto", "Trimestre"]:
                    df[col] = pd.to_numeric(df[col], errors='coerce')
            
            chart_type = st.radio("Tipo de gráfico", ["Barras", "Líneas", "Combinado"])
            
            valor_column = st.selectbox("Selecciona columna de valor", 
                                       [col for col in df.columns if col.startswith("Valor_")])
            
            if chart_type == "Barras":
                fig = px.bar(df, x="Trimestre", y=valor_column, title=f"{valor_column} por Trimestre")
            elif chart_type == "Líneas":
                fig = px.line(df, x="Trimestre", y=valor_column, title=f"{valor_column} a lo largo de los Trimestres")
            else:
                fig = go.Figure()
                fig.add_trace(go.Bar(x=df["Trimestre"], y=df[valor_column], name="Valor"))
                fig.add_trace(go.Line(x=df["Trimestre"], y=df[valor_column], name="Tendencia"))
                fig.update_layout(title=f"{valor_column} por Trimestre (Combinado)")
            
            st.plotly_chart(fig)
            
            # Mostrar tabla de datos para referencia
            st.subheader("Datos de la gráfica")
            st.dataframe(df[["Trimestre", valor_column]])
            
        except Exception as e:
            st.error(f"Error al crear el gráfico: {str(e)}")
            st.write("Datos en bruto:", all_sheets_data[sheet])
    else:
        st.warning(f"No hay datos disponibles para la hoja {sheet}")

def offer_download(zip_path):
    """
    Ofrece la descarga del archivo ZIP con todos los resultados.
    """
    st.header("Descargar Resultados")
    
    if zip_path and os.path.exists(zip_path):
        with open(zip_path, "rb") as f:
            bytes_data = f.read()
            
        st.download_button(
            label="📥 Descargar todos los resultados (ZIP)",
            data=bytes_data,
            file_name="Resultados_AnalizadorEstadisticoJudicial.zip",
            mime="application/zip"
        )
        
        st.success("El archivo ZIP contiene:")
        st.markdown("""
        - Archivos de resultados individuales para cada trimestre
        - Archivo consolidado con todos los resultados
        - Archivo de log con detalles de la ejecución
        """)
    else:
        st.warning("El archivo ZIP no está disponible. Por favor, procesa los archivos primero.")

def main():
    """
    Función principal que ejecuta la aplicación.
    """
    st.title("AnalizadorEstadisticoJudicial 📊")

    st.write("""
        [![ver código fuente](https://img.shields.io/badge/Repositorio%20GitHub-gris?logo=github)](https://github.com/bladealex9848/AnalizadorEstadisticoJudicial)
        ![Visitantes](https://api.visitorbadge.io/api/visitors?path=https%3A%2F%2Fjudidata.streamlit.app&label=Visitantes&labelColor=%235d5d5d&countColor=%231e7ebf&style=flat)
        """)

    st.markdown("""
    Esta aplicación analiza archivos Excel trimestrales y genera informes consolidados 
    para el Consejo Seccional de la Judicatura. La aplicación procesa automáticamente 
    archivos Excel con nombres específicos (Primer Trimestre.xls, Segundo Trimestre.xls, etc.) 
    y genera resultados detallados y un informe consolidado.
    """)

    # Inicializar variables de estado de la sesión
    if 'all_sheets_data' not in st.session_state:
        st.session_state.all_sheets_data = None
    if 'zip_path' not in st.session_state:
        st.session_state.zip_path = None
    if 'files_processed' not in st.session_state:
        st.session_state.files_processed = False
    if 'temp_dir' not in st.session_state:
        st.session_state.temp_dir = None
    if 'log_file' not in st.session_state:
        st.session_state.log_file = str(initialize_log())
    
    # Mostrar recursos en la barra lateral
    show_sidebar_resources()
    
    with st.sidebar:
        st.header("Carga de Archivos")
        uploaded_files = st.file_uploader("Carga tus archivos Excel trimestrales", 
                                          accept_multiple_files=True, type=['xls', 'xlsx'])
        
        st.markdown("### Instrucciones")
        st.info("""
        1. Carga tus archivos Excel trimestrales.
        2. Haz clic en 'Procesar Archivos' para analizar los datos.
        3. Visualiza los resultados en las pestañas correspondientes.
        4. Descarga el archivo ZIP con todos los resultados.
        """)

        if st.button("Usar Dataset de Muestra"):
            st.session_state.all_sheets_data = load_sample_dataset()
            st.session_state.files_processed = True
            st.success("Dataset de muestra cargado con éxito!")

        if st.button("Procesar Archivos"):
            if uploaded_files:
                with st.spinner('Procesando archivos...'):
                    try:
                        # Crear directorio temporal
                        temp_dir = tempfile.mkdtemp()
                        st.session_state.temp_dir = temp_dir
                        
                        # Guardar archivos cargados en directorio temporal
                        file_paths = []
                        for file in uploaded_files:
                            temp_file = os.path.join(temp_dir, file.name)
                            with open(temp_file, "wb") as f:
                                f.write(file.getvalue())
                            file_paths.append(temp_file)
                        
                        # Crear estructura de carpetas
                        subfolder = create_folder_structure(temp_dir)
                        
                        # Procesar archivos Excel
                        st.session_state.all_sheets_data = process_excel_files(
                            sorted_files(file_paths), subfolder, st.session_state.log_file
                        )
                        
                        if not st.session_state.all_sheets_data:
                            st.error("No se pudieron procesar los archivos. Verifica que contengan datos válidos.")
                            st.session_state.files_processed = False
                            return
                        
                        # Crear archivo consolidado
                        consolidated_file = create_consolidated_file(
                            st.session_state.all_sheets_data, subfolder, st.session_state.log_file
                        )
                        
                        if consolidated_file is None:
                            st.warning("No se pudo crear el archivo consolidado, pero los datos están disponibles para visualización.")
                        
                        # Crear archivo ZIP con todos los resultados
                        st.session_state.zip_path = create_zip_file(subfolder, st.session_state.log_file)
                        
                        st.session_state.files_processed = True
                        st.success('Archivos procesados con éxito! Puedes descargar los resultados en la pestaña "Descargar Informe".')
                        
                    except Exception as e:
                        st.error(f"Error al procesar los archivos: {str(e)}")
                        log_message(f"Error al procesar los archivos: {str(e)}", log_file=st.session_state.log_file)
                        st.info("Intente usar el método manual descargando el ejecutable o use el dataset de muestra.")
                        st.session_state.all_sheets_data = None
                        st.session_state.zip_path = None
                        st.session_state.files_processed = False
            else:
                st.warning("Por favor, carga archivos antes de procesar.")

    # Información de pie de página
    st.sidebar.markdown("---")
    st.sidebar.image("https://raw.githubusercontent.com/bladealex9848/AnalizadorEstadisticoJudicial/main/assets/logo.jpg", width=200)
    st.sidebar.write("<div style='text-align: center;'>Desarrollado por Alexander Oviedo Fadul</div>", unsafe_allow_html=True)
    st.sidebar.write("<div style='text-align: center;'>v.1.2.0</div>", unsafe_allow_html=True)
    st.sidebar.write("<div style='text-align: center;'><a href='https://github.com/bladealex9848'>GitHub</a> | <a href='https://alexanderoviedofadul.dev/'>Website</a> | <a href='https://www.linkedin.com/in/alexander-oviedo-fadul/'>LinkedIn</a></div>", unsafe_allow_html=True)

    # Mostrar pestañas con resultados
    if st.session_state.files_processed and st.session_state.all_sheets_data:
        tabs = st.tabs(["Resumen", "Detalles por Trimestre", "Gráficos", "Descargar Informe"])

        with tabs[0]:
            show_summary(st.session_state.all_sheets_data)

        with tabs[1]:
            show_details(st.session_state.all_sheets_data)

        with tabs[2]:
            show_charts(st.session_state.all_sheets_data)

        with tabs[3]:
            offer_download(st.session_state.zip_path)
    elif not st.session_state.files_processed:
        st.info("Carga tus archivos Excel y haz clic en 'Procesar Archivos' para comenzar.")
    else:
        st.warning("No se encontraron datos procesados. Por favor, asegúrate de cargar y procesar los archivos correctamente.")

    # Método manual
    st.markdown("""
    ---
    ### Método Manual
    Si experimenta problemas al cargar los archivos, puede descargar el ejecutable y procesar los archivos localmente.
    
    Instrucciones:
    1. Descargue el ejecutable 'AnalizadorEstadisticoJudicial v1.1.exe'.
    2. Coloque el ejecutable en la misma carpeta que sus archivos Excel trimestrales, o
       cree una carpeta 'PROCESAR' en C: y coloque allí los archivos Excel.
    3. Ejecute el programa haciendo doble clic en el ejecutable.
    4. Los resultados se guardarán en una nueva carpeta 'Consolidado'.
    """)
    
    # Verificar si el ejecutable está disponible
    ejecutable_path = Path('assets/AnalizadorEstadisticoJudicial v1.1.exe')
    if ejecutable_path.exists():
        st.markdown(get_binary_file_downloader_html(ejecutable_path, 'Ejecutable'), unsafe_allow_html=True)
    else:
        st.warning("El ejecutable no está disponible en esta versión web. Por favor, visite el repositorio de GitHub para obtenerlo.")
        st.markdown("[Descargar desde GitHub](https://github.com/bladealex9848/AnalizadorEstadisticoJudicial/releases)", unsafe_allow_html=True)

if __name__ == "__main__":
    warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
    warnings.filterwarnings('ignore', category=UserWarning, module='xlrd')
    main()